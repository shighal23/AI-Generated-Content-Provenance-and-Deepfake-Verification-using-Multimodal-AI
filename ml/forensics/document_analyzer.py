from pathlib import Path
import hashlib

import PyPDF2
from docx import Document


class DocumentAnalyzer:
    """
    Performs technical and basic forensic analysis of PDF and DOCX files.

    This is a structural/metadata analyzer.
    It is NOT a trained AI-generated document or forgery detector.
    """

    SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

    def analyze(self, document_path: str):
        document_path = Path(document_path)

        if not document_path.exists():
            raise FileNotFoundError(
                f"Document not found: {document_path}"
            )

        if not document_path.is_file():
            raise ValueError(
                f"Path is not a file: {document_path}"
            )

        extension = document_path.suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                "Unsupported document type. Supported types: PDF and DOCX."
            )

        file_size = document_path.stat().st_size
        sha256_hash = self._calculate_sha256(document_path)

        if extension == ".pdf":
            result = self._analyze_pdf(document_path)

        elif extension == ".docx":
            result = self._analyze_docx(document_path)

        else:
            raise ValueError("Unsupported document format.")

        result.update(
            {
                "method": "Document Forensic Analysis",
                "analysis_status": "COMPLETED",
                "filename": document_path.name,
                "file_type": extension.replace(".", "").upper(),
                "file_size_bytes": file_size,
                "sha256": sha256_hash,
            }
        )

        return result

    def _calculate_sha256(self, file_path: Path):
        sha256 = hashlib.sha256()

        with open(file_path, "rb") as file:
            while True:
                chunk = file.read(1024 * 1024)

                if not chunk:
                    break

                sha256.update(chunk)

        return sha256.hexdigest()

    def _analyze_pdf(self, file_path: Path):
        pages = 0
        text_length = 0
        metadata = {}
        errors = []

        try:
            with open(file_path, "rb") as file:
                reader = PyPDF2.PdfReader(file)

                pages = len(reader.pages)

                raw_metadata = reader.metadata

                if raw_metadata:
                    for key, value in raw_metadata.items():
                        clean_key = str(key).replace("/", "")
                        metadata[clean_key] = str(value)

                for page_number, page in enumerate(
                    reader.pages, start=1
                ):
                    try:
                        page_text = page.extract_text() or ""
                        text_length += len(page_text)
                    except Exception as exc:
                        errors.append(
                            f"Page {page_number}: {str(exc)}"
                        )

        except Exception as exc:
            errors.append(str(exc))

        metadata_fields = list(metadata.keys())

        has_metadata = bool(metadata)

        if errors:
            structural_status = "PARTIAL"
        else:
            structural_status = "NORMAL"

        return {
            "format": "PDF",
            "page_count": pages,
            "text_length": text_length,
            "metadata": metadata,
            "metadata_fields": metadata_fields,
            "has_metadata": has_metadata,
            "structural_status": structural_status,
            "errors": errors,
        }

    def _analyze_docx(self, file_path: Path):
        paragraph_count = 0
        non_empty_paragraph_count = 0
        table_count = 0
        inline_shape_count = 0
        text_length = 0
        metadata = {}

        errors = []

        try:
            document = Document(str(file_path))

            paragraph_count = len(document.paragraphs)

            for paragraph in document.paragraphs:
                text = paragraph.text or ""

                if text.strip():
                    non_empty_paragraph_count += 1

                text_length += len(text)

            table_count = len(document.tables)
            inline_shape_count = len(document.inline_shapes)

            properties = document.core_properties

            metadata = {
                "author": properties.author or "",
                "title": properties.title or "",
                "subject": properties.subject or "",
                "keywords": properties.keywords or "",
                "comments": properties.comments or "",
                "last_modified_by": properties.last_modified_by or "",
                "category": properties.category or "",
                "created": (
                    properties.created.isoformat()
                    if properties.created
                    else None
                ),
                "modified": (
                    properties.modified.isoformat()
                    if properties.modified
                    else None
                ),
            }

            metadata = {
                key: value
                for key, value in metadata.items()
                if value not in ("", None)
            }

        except Exception as exc:
            errors.append(str(exc))

        has_metadata = bool(metadata)

        if errors:
            structural_status = "PARTIAL"
        else:
            structural_status = "NORMAL"

        return {
            "format": "DOCX",
            "paragraph_count": paragraph_count,
            "non_empty_paragraph_count": non_empty_paragraph_count,
            "table_count": table_count,
            "embedded_image_count": inline_shape_count,
            "text_length": text_length,
            "metadata": metadata,
            "has_metadata": has_metadata,
            "structural_status": structural_status,
            "errors": errors,
        }